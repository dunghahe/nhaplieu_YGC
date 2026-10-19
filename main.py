import os
from io import BytesIO

from fastapi import FastAPI, HTTPException, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
import openpyxl

import database
from models import (
    ApiMessage,
    DeleteRequest,
    EntryUpdate,
    LotBaviaSaveInput,
    LotBaviaWorkRowInput,
    LotBlankReceiptSaveInput,
    LotQcSaveInput,
    LotWarehouseSaveInput,
    MachineCreate,
    MachineUpdate,
    OrderCreate,
    PartCreate,
    PartStepUpdateInput,
    LotPrintRangeInput,
    StepMachineLogInput,
    WarehouseReceiptInput,
    WorkerCreate,
    WorkerUpdate,
)


# FastAPI app chính: phục vụ SPA, API nội bộ và static files.
app = FastAPI(title="YGC MES Nhập liệu sản xuất")
DELETE_CODE = os.getenv("MES_DELETE_CODE", "admin123")


def verify_delete_code(payload: DeleteRequest) -> None:
    """Chặn thao tác xóa nếu người dùng không có mã xóa nội bộ."""
    if payload.delete_code.strip() != DELETE_CODE:
        raise HTTPException(status_code=403, detail="Mã xóa không đúng hoặc không có quyền xóa")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def startup() -> None:
    """Khởi tạo/migrate database mỗi khi server start."""
    database.init_db()


app.mount("/static", StaticFiles(directory="static"), name="static")


@app.get("/")
def home() -> Response:
    """Trang Single Page App."""
    with open("static/index.html", "rb") as f:
        content = f.read()
    return Response(
        content=content,
        media_type="text/html",
        headers={
            "Cache-Control": "no-cache, no-store, must-revalidate",
            "Pragma": "no-cache",
            "Expires": "0"
        }
    )


@app.get("/api/qr")
def api_qr(text: str = Query(..., min_length=1), size: int = Query(120, ge=80, le=360)) -> Response:
    """Tạo ảnh QR nội bộ để không phụ thuộc CDN/internet ở máy xưởng."""
    try:
        import qrcode
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="Chưa cài thư viện qrcode. Chạy: pip install -r requirements.txt") from exc

    qr = qrcode.QRCode(border=1, box_size=10)
    qr.add_data(text)
    qr.make(fit=True)
    image = qr.make_image(fill_color="#0f172a", back_color="white").convert("RGB")
    image = image.resize((size, size))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return Response(content=buffer.getvalue(), media_type="image/png")


@app.get("/api/parts")
def api_get_parts():
    """Danh sách mã hàng kèm công đoạn/cycle time."""
    return database.get_parts()


@app.get("/api/parts/summary")
def api_get_parts_summary(date_from: str | None = Query(None), date_to: str | None = Query(None)):
    """Báo cáo tổng hợp theo mã hàng, có lọc theo ngày tạo đơn hàng."""
    return database.get_parts_summary(date_from=date_from, date_to=date_to)


@app.post("/api/parts", status_code=201)
def api_create_part(payload: PartCreate):
    """Tạo mã hàng mới."""
    try:
        steps = [step.model_dump() for step in payload.steps]
        return database.create_part(payload.part_code, steps)
    except ValueError as exc:
        status_code = 409 if "tồn tại" in str(exc) else 422
        raise HTTPException(status_code=status_code, detail=str(exc)) from exc


@app.put("/api/parts/{part_code}")
def api_update_part(part_code: str, payload: PartCreate):
    """Cập nhật mã hàng và danh sách công đoạn."""
    try:
        steps = [step.model_dump() for step in payload.steps]
        return database.update_part(part_code.upper(), payload.part_code, steps)
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        status_code = 409 if "tồn tại" in str(exc) else 422
        raise HTTPException(status_code=status_code, detail=str(exc)) from exc


@app.delete("/api/parts/{part_code}")
def api_delete_part(part_code: str, payload: DeleteRequest):
    """Xóa mã hàng (phải khớp delete_code)."""
    verify_delete_code(payload)
    try:
        database.delete_part(part_code.upper())
        return {"message": "Đã xóa mã hàng"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.put("/api/parts/{part_code}/steps")
def api_update_part_steps(part_code: str, steps: list[PartStepUpdateInput]):
    """Cập nhật lại quy trình công đoạn cho một mã hàng."""
    try:
        steps_dump = [s.model_dump() for s in steps]
        database.update_part_steps(part_code.upper(), steps_dump)
        return {"message": "Đã cập nhật công đoạn thành công"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/orders")
def api_get_orders():
    """Danh sách đơn hàng kèm tiến độ."""
    return database.get_orders()


@app.get("/api/orders/{order_code}")
def api_get_order(order_code: str):
    """Lay chi tiet mot don hang de fill form va in bieu lo."""
    order = database.get_order_with_lots(order_code.upper())
    if not order:
        raise HTTPException(status_code=404, detail="Khong tim thay don hang")
    return order


@app.post("/api/orders/{order_code}/print-range")
def api_mark_order_print_range(order_code: str, payload: LotPrintRangeInput):
    """Đánh dấu một dải lô trong đơn là đã in."""
    try:
        return database.mark_order_lots_printed(order_code.upper(), payload.start_index, payload.end_index)
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.put("/api/lots/{lot_code}/blank-receipt")
def api_save_lot_blank_receipt(lot_code: str, payload: LotBlankReceiptSaveInput):
    """Lưu thông tin nhận phôi cho một lô."""
    try:
        return database.save_lot_blank_receipt(
            lot_code.upper(),
            payload.received_by,
            payload.received_at,
        )
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
def order_extra_payload(payload: OrderCreate) -> dict:
    """Tách các trường Phần I để lưu vào đơn hàng."""
    data = payload.model_dump()
    return {field: data.get(field) for field in database.ORDER_EXTRA_FIELDS}


@app.post("/api/orders", status_code=201)
def api_create_order(payload: OrderCreate):
    """Tạo đơn hàng và sinh lô."""
    try:
        return database.create_order(
            payload.order_code,
            payload.part_code,
            payload.total_qty,
            payload.qty_per_lot,
            order_extra_payload(payload),
        )
    except ValueError as exc:
        status_code = 409 if "tồn tại" in str(exc) else 422
        raise HTTPException(status_code=status_code, detail=str(exc)) from exc
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.put("/api/orders/{order_code}")
def api_update_order(order_code: str, payload: OrderCreate):
    """Cập nhật dữ liệu đơn hàng; bảo vệ đơn đã có nhập liệu khỏi chia lại lô."""
    if payload.order_code != order_code.upper():
        raise HTTPException(status_code=422, detail="Không được đổi mã PO sau khi đã tạo")
    try:
        return database.update_order(
            order_code.upper(),
            payload.part_code,
            payload.total_qty,
            payload.qty_per_lot,
            order_extra_payload(payload),
        )
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc).replace("SL/lô", "số lô")) from exc


@app.delete("/api/orders/{order_code}")
def api_delete_order(order_code: str, payload: DeleteRequest):
    """Xóa đơn hàng, lots và dữ liệu nhập liệu bên trong sau khi xác thực mã xóa."""
    verify_delete_code(payload)
    try:
        database.delete_order(order_code.upper())
        return {"message": "Đã xóa đơn hàng và dữ liệu lô bên trong"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/machines")
def api_get_machines():
    """Danh sách máy CNC cho master data và dropdown nhập liệu."""
    return database.get_machines(include_inactive=True)


@app.post("/api/machines", status_code=201)
def api_create_machine(payload: MachineCreate):
    """Tạo máy CNC mới."""
    try:
        return database.create_machine(payload.machine_code, payload.machine_name, payload.location)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@app.put("/api/machines/{machine_code}")
def api_update_machine(machine_code: str, payload: MachineUpdate):
    """Cập nhật thông tin máy CNC."""
    try:
        return database.update_machine(
            machine_code,
            payload.machine_name,
            payload.location,
            payload.is_active,
        )
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/machines/template")
def api_machine_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Mã máy", "Tên máy", "Vị trí"])
    ws.append(["M-01", "Máy Phay 01", "Xưởng 1"])
    buffer = BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Machine_Template.xlsx"}
    )


@app.post("/api/machines/import")
def api_import_machines(file: UploadFile = File(...)):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ file Excel (.xlsx)")
    
    wb = openpyxl.load_workbook(BytesIO(file.file.read()), data_only=True)
    sheet = wb.active
    
    imported = 0
    errors = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]: continue
        machine_code = str(row[0]).strip()
        machine_name = str(row[1]).strip() if len(row) > 1 and row[1] else machine_code
        location = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        
        try:
            database.create_machine(machine_code, machine_name, location)
            imported += 1
        except ValueError as e:
            errors.append(f"Dòng {row_idx} ({machine_code}): {e}")
            
    return {"message": f"Đã nhập {imported} máy", "errors": errors}


@app.delete("/api/machines/{machine_code}")
def api_delete_machine(machine_code: str, payload: DeleteRequest):
    """Soft delete máy CNC bằng cách set inactive (phải khớp delete_code)."""
    verify_delete_code(payload)
    try:
        database.set_machine_active(machine_code, 0)
        return {"message": "Đã vô hiệu máy"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/workers")
def api_get_workers():
    """Danh sách nhân công cho master data và dropdown nhập liệu."""
    return database.get_workers(include_inactive=True)


@app.post("/api/workers", status_code=201)
def api_create_worker(payload: WorkerCreate):
    """Tạo nhân công mới."""
    try:
        return database.create_worker(payload.worker_code, payload.worker_name)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@app.put("/api/workers/{worker_code}")
def api_update_worker(worker_code: str, payload: WorkerUpdate):
    """Cập nhật nhân công."""
    try:
        return database.update_worker(worker_code, payload.worker_name, payload.is_active)
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/workers/template")
def api_worker_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Mã NV", "Tên nhân công"])
    ws.append(["NV-01", "Nguyễn Văn A"])
    buffer = BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Worker_Template.xlsx"}
    )


@app.post("/api/workers/import")
def api_import_workers(file: UploadFile = File(...)):
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ file Excel (.xlsx)")
    
    wb = openpyxl.load_workbook(BytesIO(file.file.read()), data_only=True)
    sheet = wb.active
    
    imported = 0
    errors = []
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]: continue
        worker_code = str(row[0]).strip()
        worker_name = str(row[1]).strip() if len(row) > 1 and row[1] else worker_code
        
        try:
            database.create_worker(worker_code, worker_name)
            imported += 1
        except ValueError as e:
            errors.append(f"Dòng {row_idx} ({worker_code}): {e}")
            
    return {"message": f"Đã nhập {imported} nhân công", "errors": errors}


@app.delete("/api/workers/{worker_code}")
def api_delete_worker(worker_code: str, payload: DeleteRequest):
    """Soft delete nhân công bằng cách set inactive (phải khớp delete_code)."""
    verify_delete_code(payload)
    try:
        database.set_worker_active(worker_code, 0)
        return {"message": "Đã vô hiệu nhân công"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/lots/{lot_code}")
def api_get_lot(lot_code: str):
    """Thông tin lô và các công đoạn đã khởi tạo."""
    lot = database.get_lot(lot_code.upper())
    if not lot:
        raise HTTPException(status_code=404, detail="Không tìm thấy lô")
    return lot


@app.get("/api/lots/{lot_code}/qc")
def api_get_lot_qc(lot_code: str):
    """Lấy form QC của lô, kèm số lượng max theo công đoạn cuối."""
    qc = database.get_lot_qc(lot_code.upper())
    if not qc:
        raise HTTPException(status_code=404, detail="Không tìm thấy lô")
    return qc


@app.put("/api/lots/{lot_code}/qc")
def api_save_lot_qc(lot_code: str, payload: LotQcSaveInput):
    """Lưu form QC của lô."""
    try:
        return database.save_lot_qc(lot_code.upper(), payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/lots/{lot_code}/bavia")
def api_get_lot_bavia(lot_code: str):
    """Lấy form Bavia của lô, kèm số lượng max theo công đoạn cuối."""
    bavia = database.get_lot_bavia(lot_code.upper())
    if not bavia:
        raise HTTPException(status_code=404, detail="Không tìm thấy lô")
    return bavia


@app.put("/api/lots/{lot_code}/bavia")
def api_save_lot_bavia(lot_code: str, payload: LotBaviaSaveInput):
    """Lưu form Bavia của lô."""
    try:
        return database.save_lot_bavia(lot_code.upper(), payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/lots/{lot_code}/bavia-work")
def api_get_lot_bavia_work(lot_code: str):
    """Lấy form nhập sản lượng Bavia của lô."""
    work = database.get_lot_bavia_work(lot_code.upper())
    if not work:
        raise HTTPException(status_code=404, detail="Không tìm thấy lô")
    return work


@app.post("/api/lots/{lot_code}/bavia-work/rows", status_code=201)
def api_create_lot_bavia_work_row(lot_code: str, payload: LotBaviaWorkRowInput):
    """Thêm dòng máy/chuyển máy cho Bavia của lô."""
    try:
        return database.create_bavia_work_row(lot_code.upper(), payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.put("/api/lots/{lot_code}/bavia-work/rows/{line_index}")
def api_update_lot_bavia_work_row(lot_code: str, line_index: int, payload: LotBaviaWorkRowInput):
    """Cập nhật một dòng máy của Bavia."""
    try:
        return database.update_bavia_work_row(lot_code.upper(), line_index, payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.delete("/api/lots/{lot_code}/bavia-work/rows/{line_index}")
def api_delete_lot_bavia_work_row(lot_code: str, line_index: int):
    """Xóa một dòng máy của Bavia."""
    try:
        database.delete_bavia_work_row(lot_code.upper(), line_index)
        return {"message": "Đã xóa dòng máy"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/lots/{lot_code}/warehouse")
def api_get_lot_warehouse(lot_code: str):
    """Lấy form nhập kho thành phẩm của lô."""
    warehouse = database.get_lot_warehouse(lot_code.upper())
    if not warehouse:
        raise HTTPException(status_code=404, detail="Không tìm thấy lô")
    return warehouse


@app.put("/api/lots/{lot_code}/warehouse")
def api_save_lot_warehouse(lot_code: str, payload: LotWarehouseSaveInput):
    """Lưu form nhập kho thành phẩm của lô."""
    try:
        return database.save_lot_warehouse(lot_code.upper(), payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.post("/api/lots/{lot_code}/init-steps")
def api_init_lot_steps(lot_code: str):
    """Khởi tạo công đoạn cho lô khi quét lần đầu."""
    try:
        return database.init_lot_steps(lot_code.upper())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/entries/{form_id}")
def api_get_entry(form_id: str):
    """Lấy dữ liệu một biểu mẫu công đoạn theo form_id."""
    entry = database.get_entry(form_id.upper())
    if not entry:
        raise HTTPException(status_code=404, detail="Không tìm thấy biểu mẫu")
    return entry


@app.put("/api/entries/{form_id}")
def api_update_entry(form_id: str, payload: EntryUpdate):
    """Cập nhật dữ liệu legacy/ngày/kíp của một biểu mẫu công đoạn."""
    try:
        return database.update_entry(form_id.upper(), payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/entries/{form_id}/machines")
def api_get_entry_machines(form_id: str):
    """Lấy bảng máy thực hiện của một form."""
    return database.get_machine_logs(form_id.upper())


@app.post("/api/entries/{form_id}/machines", status_code=201)
def api_create_entry_machine(form_id: str, payload: StepMachineLogInput):
    """Thêm dòng máy/chuyển máy cho form."""
    try:
        return database.create_machine_log(form_id.upper(), payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.put("/api/entries/{form_id}/machines/{line_index}")
def api_update_entry_machine(form_id: str, line_index: int, payload: StepMachineLogInput):
    """Cập nhật một dòng máy của form."""
    try:
        return database.update_machine_log(form_id.upper(), line_index, payload.model_dump())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.delete("/api/entries/{form_id}/machines/{line_index}")
def api_delete_entry_machine(form_id: str, line_index: int):
    """Xóa một dòng máy của form."""
    try:
        database.delete_machine_log(form_id.upper(), line_index)
        return {"message": "Đã xóa dòng máy"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/reports/summary")
def api_summary_report(date_from: str | None = Query(None), date_to: str | None = Query(None)):
    """Báo cáo tổng quan theo đơn hàng."""
    return database.get_summary_report(date_from, date_to)


@app.get("/api/reports/order/{order_code}")
def api_order_report(order_code: str, date_from: str | None = Query(None), date_to: str | None = Query(None)):
    """Báo cáo chi tiết từng lô/công đoạn trong đơn hàng."""
    return database.get_order_report(order_code.upper(), date_from, date_to)


@app.get("/api/reports/order/{order_code}/quality")
def api_order_quality_report(order_code: str, date_from: str | None = Query(None), date_to: str | None = Query(None)):
    """Dữ liệu QC, Bavia và kiểm tra kho trong đơn hàng."""
    try:
        return database.get_order_quality_report(order_code.upper(), date_from, date_to)
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/reports/machines")
def api_machine_report(date_from: str | None = Query(None), date_to: str | None = Query(None)):
    """Báo cáo hiệu suất máy."""
    return database.get_machine_efficiency_report(date_from, date_to)


# ─── API Kho thành phẩm mới ──────────────────────────────────────────────────

@app.get("/api/warehouse/orders")
def api_warehouse_orders(search: str = Query("")):
    """Danh sách đơn hàng kèm tiến độ nhập kho (trang chủ kho)."""
    return database.get_warehouse_orders_summary(search.strip())


@app.get("/api/warehouse/orders/{order_code}/detail")
def api_warehouse_order_detail(order_code: str):
    """Chi tiết đơn kho thành phẩm kèm danh sách lô."""
    try:
        return database.get_warehouse_order_detail(order_code.upper())
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/api/lots/{lot_code}/warehouse-receipts", status_code=201)
def api_add_warehouse_receipt(lot_code: str, payload: WarehouseReceiptInput):
    """Thêm phiếu nhập kho cho lô."""
    try:
        return database.add_warehouse_receipt(
            lot_code.upper(),
            payload.qty,
            str(payload.receipt_date or ""),
            payload.note,
        )
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.delete("/api/lots/{lot_code}/warehouse-receipts/{receipt_id}")
def api_delete_warehouse_receipt(lot_code: str, receipt_id: int, payload: DeleteRequest):
    """Xóa phiếu nhập kho."""
    verify_delete_code(payload)
    try:
        database.delete_warehouse_receipt(lot_code.upper(), receipt_id)
        return {"message": "Đã xóa phiếu nhập kho"}
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
